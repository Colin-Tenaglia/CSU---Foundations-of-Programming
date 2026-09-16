"""Load a delimited or spreadsheet file into a plain list of string rows.

QuickBooks writes CSV with a BOM, TSV from some report screens, and users
frequently hand us the .xlsx they saved out of Excel instead. All three land
here and come out looking the same.
"""

from __future__ import annotations

import csv
import os
from typing import List

csv.field_size_limit(10_000_000)


def _sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        # Sniffer fails on reports whose first rows are a single title cell.
        counts = {d: sample.count(d) for d in ",\t;|"}
        best = max(counts, key=counts.get)
        return best if counts[best] else ","


def load_rows(path: str) -> List[List[str]]:
    """Return the file as rows of stripped strings."""
    extension = os.path.splitext(path)[1].lower()
    if extension in (".xlsx", ".xlsm"):
        return _load_excel(path)

    with open(path, newline="", encoding="utf-8-sig", errors="replace") as handle:
        sample = handle.read(64 * 1024)
        handle.seek(0)
        delimiter = "\t" if extension == ".iif" else _sniff_delimiter(sample)
        return [[(cell or "").strip() for cell in row] for row in csv.reader(handle, delimiter=delimiter)]


def _load_excel(path: str) -> List[List[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise SystemExit(
            f"Reading {path} needs openpyxl. Install it with `pip install openpyxl`, "
            "or re-save the file as CSV from Excel."
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: List[List[str]] = []
    for sheet_row in workbook[workbook.sheetnames[0]].iter_rows(values_only=True):
        rows.append(["" if cell is None else str(cell).strip() for cell in sheet_row])
    workbook.close()
    return rows


def is_blank(row: List[str]) -> bool:
    return not any(cell for cell in row)
